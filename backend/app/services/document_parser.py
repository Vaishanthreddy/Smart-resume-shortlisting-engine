"""Document ingestion for PDF, DOCX and TXT files.

Uploaded bytes are parsed in memory and discarded; nothing is written to disk
unless config.PERSIST_UPLOADS is explicitly turned on.
"""

from __future__ import annotations

import io
import logging
from dataclasses import dataclass, field
from pathlib import Path

from app.config import (
    ALLOWED_EXTENSIONS,
    ALLOWED_MIME_TYPES,
    MAX_FILE_SIZE_BYTES,
    MAX_FILE_SIZE_MB,
    MIN_DOCUMENT_CHARS,
    MIN_PDF_CHARS_PER_PAGE,
)
from app.services.text_cleaner import clean_text

logger = logging.getLogger(__name__)


class DocumentParseError(Exception):
    """Raised when a document cannot be turned into usable text.

    The message is user-facing and must explain what to do about it.
    """

    def __init__(self, message: str, *, reason_code: str = "parse_error") -> None:
        super().__init__(message)
        self.message = message
        self.reason_code = reason_code


@dataclass
class ExtractedDocument:
    filename: str
    text: str
    extension: str
    char_count: int
    page_count: int = 0
    warnings: list[str] = field(default_factory=list)


def _extension_of(filename: str) -> str:
    return Path(filename).suffix.lower()


def validate_upload(filename: str, data: bytes, content_type: str | None = None) -> str:
    """Validate extension, size, MIME type and emptiness. Returns the extension."""
    extension = _extension_of(filename)

    if extension not in ALLOWED_EXTENSIONS:
        allowed = ", ".join(sorted(ALLOWED_EXTENSIONS))
        raise DocumentParseError(
            f"'{filename}' has an unsupported file type. Upload {allowed} files.",
            reason_code="unsupported_extension",
        )

    if not data:
        raise DocumentParseError(
            f"'{filename}' is empty. Upload a file that contains text.",
            reason_code="empty_file",
        )

    if len(data) > MAX_FILE_SIZE_BYTES:
        actual_mb = len(data) / (1024 * 1024)
        raise DocumentParseError(
            f"'{filename}' is {actual_mb:.1f} MB, over the {MAX_FILE_SIZE_MB:.0f} MB limit.",
            reason_code="file_too_large",
        )

    # MIME validation where practical: browsers are inconsistent, so an unknown
    # or generic type is accepted and the extension decides the parser.
    if content_type:
        normalized = content_type.split(";")[0].strip().lower()
        permitted = ALLOWED_MIME_TYPES.get(extension, set())
        if normalized and permitted and normalized not in permitted:
            logger.warning(
                "Unexpected content type for %s: %s (parsing by extension)",
                extension, normalized,
            )

    return extension


def _extract_pdf(filename: str, data: bytes) -> tuple[str, int]:
    try:
        import fitz  # PyMuPDF
    except ImportError as exc:  # pragma: no cover - dependency guaranteed by requirements
        raise DocumentParseError(
            "PDF support is unavailable because PyMuPDF is not installed.",
            reason_code="missing_dependency",
        ) from exc

    try:
        with fitz.open(stream=data, filetype="pdf") as document:
            if document.is_encrypted and not document.authenticate(""):
                raise DocumentParseError(
                    f"'{filename}' is password protected. Upload an unlocked copy.",
                    reason_code="encrypted_pdf",
                )
            pages = [page.get_text("text") or "" for page in document]
            page_count = len(pages)
    except DocumentParseError:
        raise
    except Exception as exc:
        raise DocumentParseError(
            f"'{filename}' could not be read as a PDF. The file may be corrupt.",
            reason_code="corrupt_pdf",
        ) from exc

    text = "\n".join(pages)
    stripped_len = len(text.strip())

    if page_count and stripped_len < MIN_PDF_CHARS_PER_PAGE * page_count:
        raise DocumentParseError(
            f"'{filename}' appears to be a scanned or image-only PDF — no selectable "
            "text was found. Upload a text-based PDF or a DOCX/TXT version.",
            reason_code="image_only_pdf",
        )

    return text, page_count


def _extract_docx(filename: str, data: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:  # pragma: no cover
        raise DocumentParseError(
            "DOCX support is unavailable because python-docx is not installed.",
            reason_code="missing_dependency",
        ) from exc

    try:
        document = docx.Document(io.BytesIO(data))
    except Exception as exc:
        raise DocumentParseError(
            f"'{filename}' could not be read as a DOCX file. If it is an older .doc "
            "file, save it as .docx and upload again.",
            reason_code="corrupt_docx",
        ) from exc

    parts: list[str] = [paragraph.text for paragraph in document.paragraphs]

    # Many resumes lay content out in tables; read those cells too.
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(c for c in cells if c)
            if line:
                parts.append(line)

    return "\n".join(parts)


def _extract_txt(filename: str, data: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "latin-1"):
        try:
            return data.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    # Last resort: never fail a text file outright.
    return data.decode("utf-8", errors="replace")


def extract_document(
    filename: str,
    data: bytes,
    content_type: str | None = None,
) -> ExtractedDocument:
    """Validate and extract cleaned text from an uploaded file.

    Raises DocumentParseError with a user-facing message on failure.
    """
    extension = validate_upload(filename, data, content_type)
    warnings: list[str] = []
    page_count = 0

    if extension == ".pdf":
        raw_text, page_count = _extract_pdf(filename, data)
    elif extension == ".docx":
        raw_text = _extract_docx(filename, data)
    else:
        raw_text = _extract_txt(filename, data)

    text = clean_text(raw_text)

    if len(text) < MIN_DOCUMENT_CHARS:
        raise DocumentParseError(
            f"'{filename}' contains no usable text (only {len(text)} characters were "
            "extracted). Check that the file is not blank or image-only.",
            reason_code="empty_document",
        )

    if len(text) < 400:
        warnings.append(
            "Very little text was extracted; matching may be limited for this document."
        )

    logger.info(
        "Extracted %s: %d characters, %d pages", extension, len(text), page_count
    )

    return ExtractedDocument(
        filename=filename,
        text=text,
        extension=extension,
        char_count=len(text),
        page_count=page_count,
        warnings=warnings,
    )
