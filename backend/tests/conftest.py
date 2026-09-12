"""Shared fixtures.

Every document here is synthetic and written for these tests. No hackathon
evaluation document is used, and nothing is trained on any of it.
"""

from __future__ import annotations

import io
import os

# Keep the suite offline and fast: the lexical fallback vector space is
# deterministic and needs no model download.
os.environ.setdefault("FORCE_FALLBACK_EMBEDDINGS", "true")

import pytest  # noqa: E402

from app.services.keyword_matcher import get_taxonomy  # noqa: E402
from app.services.ranking_engine import UploadedFile  # noqa: E402
from app.services.semantic_matcher import get_semantic_engine  # noqa: E402

JD_TEXT = """Job Title: Backend Engineer

About us:
We are an equal opportunity employer building payment infrastructure.

Responsibilities:
Develop and maintain RESTful backend services for the payments platform.
Design database schemas and optimise slow queries.
Collaborate with the frontend team to agree API contracts.
Deploy services using containerised workflows.

Required Skills:
Proficient in Python and PostgreSQL.
Strong knowledge of REST API design.
Demonstrated experience with Docker.

Preferred Qualifications:
Familiarity with Kubernetes is a plus.
Exposure to Redis is nice to have.

Education:
Bachelor's degree in Computer Science is required.

Certifications:
AWS Certified Developer certification preferred.

Soft Skills:
Strong communication skills and teamwork are required.
"""

STRONG_RESUME = """Dana Whitfield
dana.whitfield@example.com

Summary
Backend engineer with four years building transaction services.

Technical Skills
Python, PostgreSQL, REST API, Docker, Kubernetes, Redis, Git

Work Experience
Backend Engineer, Latimer Payments (2021 - 2025)
Developed and maintained RESTful services handling card settlement traffic.
Designed normalised database schemas and tuned slow PostgreSQL queries.
Deployed every service through Docker containers and a CI pipeline.
Collaborated with a five-person frontend team to agree API contracts.

Projects
Ledger reconciliation service: built a Python service exposing REST endpoints.

Education
Bachelor of Science in Computer Science, 2021

Certifications
AWS Certified Developer - Associate, 2023
"""

MEDIUM_RESUME = """Ravi Sekhar
ravi.sekhar@example.com

Skills
Python, MySQL, Flask, HTML, CSS

Experience
Software Intern, Delta Systems (2024)
Wrote Python scripts to import customer records into MySQL.
Assisted with a small internal reporting tool.

Projects
Library management system built with Flask and SQLite.

Education
Bachelor of Engineering in Information Technology, 2024
"""

WEAK_RESUME = """Toby Marsh
toby.marsh@example.com

Skills
Photoshop, Illustrator, social media scheduling, copywriting

Experience
Graphic Designer, Rowan Studio (2020 - 2024)
Produced print and social artwork for hospitality clients.

Education
Diploma of Visual Arts, 2020
"""

# A resume with no headings at all — formatting must not cost points.
UNFORMATTED_RESUME = """kim delacroix python postgresql rest api docker kubernetes
built and maintained restful backend services for an internal payments tool
designed database schemas and optimised slow queries for reporting
deployed containerised services with docker and collaborated with a small team
bachelor of science in computer science 2020
"""


@pytest.fixture(scope="session")
def taxonomy():
    return get_taxonomy()


@pytest.fixture(scope="session")
def engine():
    return get_semantic_engine()


@pytest.fixture
def jd_text() -> str:
    return JD_TEXT


@pytest.fixture
def jd_upload() -> UploadedFile:
    return UploadedFile("jd.txt", JD_TEXT.encode("utf-8"), "text/plain")


def make_txt_upload(name: str, text: str) -> UploadedFile:
    return UploadedFile(name, text.encode("utf-8"), "text/plain")


@pytest.fixture
def resume_uploads() -> list[UploadedFile]:
    return [
        make_txt_upload("a_strong.txt", STRONG_RESUME),
        make_txt_upload("b_medium.txt", MEDIUM_RESUME),
        make_txt_upload("c_weak.txt", WEAK_RESUME),
    ]


# --------------------------------------------------------------------------- #
# Generated binary documents
# --------------------------------------------------------------------------- #
@pytest.fixture(scope="session")
def text_pdf_bytes() -> bytes:
    import fitz

    document = fitz.open()
    page = document.new_page()
    page.insert_text((72, 96), "Dana Whitfield", fontsize=14)
    page.insert_text((72, 120), "Python PostgreSQL Docker REST API", fontsize=11)
    page.insert_text((72, 144), "Developed RESTful backend services.", fontsize=11)
    page.insert_text((72, 168), "Bachelor of Science in Computer Science", fontsize=11)
    data = document.tobytes()
    document.close()
    return data


@pytest.fixture(scope="session")
def image_only_pdf_bytes() -> bytes:
    """A PDF with pages but no selectable text."""
    import fitz

    document = fitz.open()
    document.new_page()
    document.new_page()
    data = document.tobytes()
    document.close()
    return data


@pytest.fixture(scope="session")
def docx_bytes() -> bytes:
    import docx

    document = docx.Document()
    document.add_paragraph("Dana Whitfield")
    document.add_paragraph("Technical Skills")
    document.add_paragraph("Python, PostgreSQL, Docker, REST API")
    document.add_paragraph("Work Experience")
    document.add_paragraph("Developed and maintained RESTful backend services.")

    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Education"
    table.rows[0].cells[1].text = "Bachelor of Science in Computer Science"

    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()
